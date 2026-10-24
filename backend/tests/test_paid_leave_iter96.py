"""Iter96: Active-only monthly generation, Approved->Rejected paid restore, re-approve re-reserve,
Paid Leave Balances Excel export. LOCAL DB only, cleans all test data."""
import io
import os
import subprocess
import pytest
import requests
import openpyxl
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
db = MongoClient("mongodb://localhost:27017")["test_database"]
SEED = "/app/backend/tests/seed_paid_leave_test_data.py"
REASON = "TEST_PL_ iter96 reason text"
GEN_Y, GEN_M = 2030, 2
_tokens = {}
INIT = {"fulltime2": 2.0, "fulltime0": 0.0, "halfbal": 0.5, "intern": 0.0, "advzero": 0.0,
        "unconfirmed": 0.0, "inactive": 1.0, "negative": -1.0}


def _revert_generation():
    if db.paid_leave_generation.find_one({"year": GEN_Y, "month": GEN_M}):
        for h in db.paid_leave_history.find({"action_type": "MONTHLY_GENERATION", "year": GEN_Y, "month": GEN_M}):
            db.employees.update_one({"id": h["employee_id"]}, {"$inc": {"available_paid_leave": -1}})
        db.paid_leave_history.delete_many({"action_type": "MONTHLY_GENERATION", "year": GEN_Y, "month": GEN_M})
    db.paid_leave_generation.delete_many({"year": GEN_Y, "month": GEN_M})


@pytest.fixture(scope="module", autouse=True)
def seed():
    _revert_generation()
    subprocess.run(["python3", SEED, "seed"], check=True)
    yield
    _revert_generation()
    subprocess.run(["python3", SEED, "cleanup"], check=True)
    db.notifications.delete_many({"message": {"$regex": "TEST_PL_"}})


def _emp(k):
    return db.employees.find_one({"full_name": f"TEST_PL_{k}"}, {"_id": 0})


def _bal(k):
    return round(float(_emp(k).get("available_paid_leave") or 0.0), 1)


def reset(k):
    e = _emp(k)
    db.leaves.delete_many({"employee_id": e["id"]})
    db.employees.update_one({"id": e["id"]}, {"$set": {"available_paid_leave": INIT[k]},
                                             "$unset": {"paid_leave_advance_month": "", "paid_leave_advance_used": ""}})


def tok(u, p="Test@123"):
    if u not in _tokens:
        r = requests.post(f"{API}/auth/login", json={"username": u, "password": p})
        assert r.status_code == 200, r.text
        _tokens[u] = r.json().get("token") or r.json().get("access_token")
    return {"Authorization": f"Bearer {_tokens[u]}"}


def E(k):
    return tok(f"test_pl_{k}")


def ADM():
    return tok("admin", "pass123")


def apply(k, date, split="Full Day", paid=True):
    return requests.post(f"{API}/employee/leaves/apply", headers=E(k), json={
        "leave_type": "Emergency", "leave_split": split, "start_date": date, "end_date": date,
        "reason": REASON, "consider_as_paid_leave": paid})


def lid_of(r):
    j = r.json()
    return j.get("id") or j.get("leave_id")


def lv(lid):
    return db.leaves.find_one({"id": lid}, {"_id": 0})


def approve(lid):
    return requests.put(f"{API}/leaves/{lid}/approve", headers=ADM(), json={"leave_validity": "valid"})


def reject(lid):
    return requests.put(f"{API}/leaves/{lid}/reject", headers=ADM(), json={})


# ---- Generation (active-only) ----
def test_generation_active_only():
    st = requests.get(f"{API}/admin/paid-leave/status", params={"year": GEN_Y, "month": GEN_M}, headers=ADM())
    assert st.status_code == 200, st.text
    elig = st.json()["eligible_count"]
    rows = requests.get(f"{API}/admin/paid-leave/balances", headers=ADM()).json()
    assert elig == len(rows)
    names = {r["emp_name"] for r in rows}
    assert "TEST_PL_inactive" not in names and "TEST_PL_unconfirmed" not in names and "TEST_PL_intern" not in names
    assert "TEST_PL_fulltime2" in names and "TEST_PL_negative" in names
    r = requests.post(f"{API}/admin/paid-leave/generate", headers=ADM(), json={"year": GEN_Y, "month": GEN_M})
    assert r.status_code == 200, r.text
    assert r.json()["count"] == elig
    assert str(elig) in r.json()["message"]
    assert _bal("inactive") == 1.0 and _bal("unconfirmed") == 0.0 and _bal("intern") == 0.0
    assert _bal("negative") == 0.0 and _bal("fulltime2") == 3.0
    r2 = requests.post(f"{API}/admin/paid-leave/generate", headers=ADM(), json={"year": GEN_Y, "month": GEN_M})
    assert r2.status_code == 400 and "already" in r2.text
    _revert_generation()
    assert _bal("fulltime2") == 2.0 and _bal("negative") == -1.0


# ---- Approved -> Rejected restore ----
def test_approved_full_day_reject_restores_once():
    reset("fulltime2")
    r = apply("fulltime2", "2026-10-27")
    assert r.status_code in (200, 201), r.text
    lid = lid_of(r)
    assert _bal("fulltime2") == 1.0
    assert approve(lid).status_code == 200 and _bal("fulltime2") == 1.0
    rj = reject(lid)
    assert rj.status_code == 200, rj.text
    assert rj.json()["status"] == "rejected" and rj.json()["paid_leave_reservation_status"] == "released"
    assert _bal("fulltime2") == 2.0
    assert reject(lid).status_code == 200 and _bal("fulltime2") == 2.0
    # consistency
    eb = requests.get(f"{API}/employee/paid-leave-balance", headers=E("fulltime2")).json()["balance"]
    ab = requests.get(f"{API}/admin/employees/{_emp('fulltime2')['id']}/paid-leave-balance", headers=ADM()).json()["balance"]
    tb = [x for x in requests.get(f"{API}/admin/paid-leave/balances", headers=ADM()).json()
          if x["emp_name"] == "TEST_PL_fulltime2"][0]["available_paid_leave"]
    assert eb == ab == tb == 2.0
    # re-approve re-deducts
    ra = approve(lid)
    assert ra.status_code == 200, ra.text
    assert _bal("fulltime2") == 1.0
    assert ra.json()["paid_leave_reservation_status"] == "finalized"
    assert reject(lid).status_code == 200 and _bal("fulltime2") == 2.0
    # loop check: approve/reject cycles never inflate
    for _ in range(2):
        approve(lid); reject(lid)
    assert _bal("fulltime2") == 2.0


@pytest.mark.parametrize("split", ["First Half", "Second Half"])
def test_approved_half_day_reject_restores_half(split):
    reset("fulltime2")
    r = apply("fulltime2", "2026-10-28", split=split)
    assert r.status_code in (200, 201), r.text
    lid = lid_of(r)
    assert _bal("fulltime2") == 1.5
    assert approve(lid).status_code == 200
    reject(lid)
    assert _bal("fulltime2") == 2.0
    reject(lid)
    assert _bal("fulltime2") == 2.0


def test_admin_multi_day_paid_reject_restores_two():
    reset("fulltime2")
    r = requests.post(f"{API}/leaves", headers=ADM(), json={
        "employee_id": _emp("fulltime2")["id"], "leave_type": "Casual Leave", "start_date": "2026-10-29",
        "end_date": "2026-10-30", "leave_split": "Full Day", "reason": REASON, "consider_as_paid_leave": True})
    assert r.status_code in (200, 201), r.text
    lid = lid_of(r)
    assert _bal("fulltime2") == 0.0
    a = approve(lid)
    assert a.status_code == 200, a.text
    reject(lid)
    assert _bal("fulltime2") == 2.0
    reject(lid)
    assert _bal("fulltime2") == 2.0


def test_reapprove_insufficient_400():
    reset("fulltime2")
    lid = lid_of(apply("fulltime2", "2026-10-27"))
    approve(lid); reject(lid)
    assert _bal("fulltime2") == 2.0
    db.employees.update_one({"id": _emp("fulltime2")["id"]}, {"$set": {"available_paid_leave": 0.5}})
    ra = approve(lid)
    assert ra.status_code == 400, ra.text
    assert _bal("fulltime2") == 0.5
    lv = db.leaves.find_one({"id": lid})
    assert lv["status"] == "rejected" and lv["paid_leave_reservation_status"] == "released"


def test_pending_paid_reject_restores_once():
    reset("fulltime2")
    lid = lid_of(apply("fulltime2", "2026-10-27"))
    assert _bal("fulltime2") == 1.0
    reject(lid)
    assert _bal("fulltime2") == 2.0
    reject(lid)
    assert _bal("fulltime2") == 2.0


def test_normal_leave_reject_unchanged():
    reset("fulltime2")
    r = apply("fulltime2", "2026-10-27", paid=False)
    assert r.status_code in (200, 201), r.text
    lid = lid_of(r)
    assert _bal("fulltime2") == 2.0
    rj = reject(lid)
    assert rj.status_code == 200 and rj.json()["status"] == "rejected"
    assert not rj.json().get("consider_as_paid_leave")
    assert rj.json().get("paid_leave_reservation_status") in (None, "")
    assert _bal("fulltime2") == 2.0


def test_advance_approved_reject_reverses():
    reset("fulltime0")
    r = apply("fulltime0", "2026-11-10")
    assert r.status_code in (200, 201), r.text
    lid = lid_of(r)
    assert lv(lid).get("is_advance_paid_leave") is True
    assert _bal("fulltime0") == -1.0
    assert approve(lid).status_code == 200
    rj = reject(lid)
    assert rj.status_code == 200
    assert rj.json().get("advance_settlement") == "reversed"
    assert _bal("fulltime0") == 0.0
    reject(lid)
    assert _bal("fulltime0") == 0.0


# ---- Excel export ----
def _xlsx(params=None, headers=None):
    return requests.get(f"{API}/admin/paid-leave/balances/export", params=params or {}, headers=headers or ADM())


def test_export_xlsx_contents():
    for k in INIT:
        reset(k)
    r = _xlsx()
    assert r.status_code == 200, r.text
    cd = r.headers.get("content-disposition", "")
    assert "Employee_Paid_Leave_Balances_" in cd and cd.strip().endswith(".xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == ["Employee Name", "Team", "Available Paid Leave"]
    data = {row[0]: row for row in rows[1:]}
    api_rows = requests.get(f"{API}/admin/paid-leave/balances", headers=ADM()).json()
    assert len(rows) - 1 == len(api_rows)
    for bad in ("TEST_PL_inactive", "TEST_PL_unconfirmed", "TEST_PL_intern"):
        assert bad not in data
    assert data["TEST_PL_negative"][2] == -1 and isinstance(data["TEST_PL_negative"][2], (int, float))
    assert data["TEST_PL_halfbal"][2] == 0.5
    assert data["TEST_PL_fulltime0"][2] == 0 and data["TEST_PL_fulltime0"][2] is not None
    assert data["TEST_PL_fulltime2"][1] == "QA"


def test_export_xlsx_search():
    ws = openpyxl.load_workbook(io.BytesIO(_xlsx({"search": "TEST_PL_full"}).content)).active
    names = sorted(r[0] for r in list(ws.iter_rows(values_only=True))[1:])
    assert names == ["TEST_PL_fulltime0", "TEST_PL_fulltime2"]


def test_export_xlsx_employee_forbidden():
    assert _xlsx(headers=E("fulltime2")).status_code == 403
    assert requests.get(f"{API}/admin/paid-leave/balances/export").status_code in (401, 403)
