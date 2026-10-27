# BRSF N09 "Monthly Late Arrivals (More than 3)" — API, boundaries, recalc, override, export/import
import csv
import io
import os
import subprocess
import uuid

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    with open("/app/frontend/.env") as f:
        for ln in f:
            if ln.startswith("REACT_APP_BACKEND_URL="):
                BASE_URL = ln.split("=", 1)[1].strip().rstrip("/")
MONTH = "2026-09"
TAG = "TEST_N09_"
mdb = MongoClient("mongodb://localhost:27017")["test_database"]
SEED = "/app/backend/tests/seed_brsf_late_test_data.py"
STATE = {"batches": []}


def _run_seed(cmd):
    subprocess.run(["python3", SEED, cmd], cwd="/app/backend", check=True, capture_output=True)


@pytest.fixture(scope="module", autouse=True)
def seeded():
    _run_seed("seed")
    yield
    mdb.brsf_import_batches.delete_many({"id": {"$in": STATE["batches"]}})
    _run_seed("cleanup")


@pytest.fixture(scope="module")
def s():
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": "admin", "password": "pass123"})
    if r.status_code != 200:
        pytest.skip("admin login failed")
    sess = requests.Session()
    sess.headers["Authorization"] = f"Bearer {r.json()['token']}"
    return sess


def _eid(key):
    return mdb.employees.find_one({"full_name": f"{TAG}{key}"})["id"]


def _stars(s, key):
    r = s.get(f"{BASE_URL}/api/brsf/stars", params={"employee_id": _eid(key), "month": MONTH})
    assert r.status_code == 200, r.text
    return r.json()


def _n09(data):
    return next(l for l in data["lines"] if l["code"] == "N09")


def _others(data):
    return {l["code"]: l["final_value"] for l in data["lines"] if l["code"] != "N09"}


# ---------- core seeded scenarios
def test_example_n09_minus_one(s):
    d = _stars(s, "example")
    assert len(d["lines"]) == 15
    n = _n09(d)
    assert n["name"] == "Monthly Late Arrivals (More than 3)"
    assert n["system_value"] == -1 and n["final_value"] == -1
    ch = {c["date"]: c for c in n["system_children"]}
    assert set(ch) == {"2026-09-03", "2026-09-07", "2026-09-10", "2026-09-14",
                       "2026-09-21", "2026-09-24", "2026-09-28"}
    counted = sorted(k for k, c in ch.items() if c["counted"])
    assert counted == ["2026-09-07", "2026-09-10", "2026-09-21", "2026-09-28"]
    assert ch["2026-09-07"]["reporting_time"] == "10:00"
    assert ch["2026-09-21"]["reporting_time"] == "10:30"
    assert ch["2026-09-10"]["punch_in"] == "10:15"
    for c in ch.values():
        assert {"date", "reporting_time", "punch_in", "reason"} <= set(c)
    STATE["example_others"] = _others(d)


def test_three_n09_zero(s):
    d = _stars(s, "three")
    n = _n09(d)
    assert n["system_value"] == 0 and n["final_value"] == 0
    counted = sorted(c["date"] for c in n["system_children"] if c["counted"])
    assert counted == ["2026-09-08", "2026-09-09", "2026-09-22"]
    STATE["three_others"] = _others(d)


# ---------- boundary cases on a dedicated employee
def _att(eid, ymd, hhmm):
    y, m, d = ymd.split("-")
    mdb.attendance.insert_one({"id": str(uuid.uuid4()), "employee_id": eid, "date": f"{d}-{m}-{y}",
                               "check_in_24h": hhmm, "status": "Late Login"})


def _lr(eid, ymd, created, status="pending"):
    mdb.late_requests.insert_one({"id": str(uuid.uuid4()), "employee_id": eid, "date": ymd,
                                  "status": status, "created_at": created, "reason": "t"})


def _lv(eid, ymd, split, status):
    mdb.leaves.insert_one({"id": str(uuid.uuid4()), "employee_id": eid, "leave_type": "Sick",
                           "leave_split": split, "start_date": ymd, "end_date": ymd, "status": status})


def test_boundaries(s):
    ex = mdb.employees.find_one({"full_name": f"{TAG}example"}, {"_id": 0})
    eid = str(uuid.uuid4())
    mdb.employees.insert_one({**ex, "id": eid, "full_name": f"{TAG}bnd", "custom_employee_id": f"{TAG}bnd",
                              "official_email": "bnd@example.invalid"})
    _att(eid, "2026-09-01", "10:05"); _lr(eid, "2026-09-01", "2026-09-01T09:30:00+05:30")      # counted
    _att(eid, "2026-09-02", "10:05"); _lr(eid, "2026-09-02", "2026-09-02T09:29:00+05:30")      # excluded
    _att(eid, "2026-09-04", "10:05"); _lr(eid, "2026-09-05", "2026-09-04T08:00:00+05:30")      # counted
    _att(eid, "2026-09-08", "10:05"); _lv(eid, "2026-09-08", "Second Half", "approved")        # counted
    _att(eid, "2026-09-09", "10:05"); _lv(eid, "2026-09-09", "First Half", "rejected")         # counted
    _att(eid, "2026-09-10", "10:05"); _lv(eid, "2026-09-10", "First Half", "cancelled")        # counted
    _att(eid, "2026-09-11", "10:05"); _lr(eid, "2026-09-11", "2026-09-10T20:00:00+05:30", "rejected")  # counted
    _att(eid, "2026-09-18", "10:30")                                                            # not late
    _att(eid, "2026-09-19", "10:31")                                                            # counted
    # multiple attendance docs per date impossible: unique index (employee_id, date) on attendance
    _att(eid, "2026-09-24", "10:45"); _lr(eid, "2026-09-24", "2026-09-24T03:59:00+00:00")       # UTC 09:29 IST excl
    _att(eid, "2026-09-25", "10:45"); _lr(eid, "2026-09-25", "2026-09-25T04:00:00Z")            # UTC 09:30 IST cnt
    d = _stars(s, "bnd")
    n = _n09(d)
    ch = {c["date"]: c for c in n["system_children"]}
    print({k: (c["punch_in"], c["counted"], c["reason"]) for k, c in ch.items()})
    assert ch["2026-09-01"]["counted"] is True
    assert ch["2026-09-02"]["counted"] is False
    assert ch["2026-09-04"]["counted"] is True
    assert ch["2026-09-08"]["counted"] is True
    assert ch["2026-09-09"]["counted"] is True
    assert ch["2026-09-10"]["counted"] is True
    assert ch["2026-09-11"]["counted"] is True
    assert "2026-09-18" not in ch
    assert ch["2026-09-19"]["counted"] is True
    assert len({c["date"] for c in n["system_children"]}) == len(n["system_children"])
    assert ch["2026-09-24"]["counted"] is False
    assert ch["2026-09-25"]["counted"] is True
    assert n["system_value"] == -1


# ---------- recalculation
def test_recalc_late_request_rejected(s):
    eid = _eid("three")
    mdb.late_requests.update_one({"employee_id": eid, "date": "2026-09-29"}, {"$set": {"status": "rejected"}})
    r = s.post(f"{BASE_URL}/api/brsf/recalculate", json={"month": MONTH})
    assert r.status_code == 200 and r.json()["success"]
    line = mdb.brsf_star_lines.find_one({"employee_id": eid, "year": 2026, "month": 9, "code": "N09"})
    assert line["final_value"] == -1
    d = _stars(s, "three")
    assert _n09(d)["final_value"] == -1
    assert _others(d) == STATE["three_others"], "other criteria changed"
    mdb.late_requests.update_one({"employee_id": eid, "date": "2026-09-29"}, {"$set": {"status": "pending"}})
    assert _n09(_stars(s, "three"))["final_value"] == 0


def test_recalc_fh_leave_rejected(s):
    eid = _eid("example")
    # turn 09-14 FH pending -> rejected, and make three's count unaffected; example goes 4 -> 5 counted
    mdb.leaves.update_one({"employee_id": eid, "start_date": "2026-09-14"}, {"$set": {"status": "rejected"}})
    n = _n09(_stars(s, "example"))
    assert next(c for c in n["system_children"] if c["date"] == "2026-09-14")["counted"] is True
    mdb.leaves.update_one({"employee_id": eid, "start_date": "2026-09-14"}, {"$set": {"status": "pending"}})
    # three: add a FH pending leave on a counted day then reject it
    t = _eid("three")
    _lv(t, "2026-09-22", "First Half", "pending")
    mdb.late_requests.update_one({"employee_id": t, "date": "2026-09-29"}, {"$set": {"status": "rejected"}})
    assert _n09(_stars(s, "three"))["final_value"] == 0   # 22 excluded, 29 counted -> 3
    mdb.leaves.update_one({"employee_id": t, "start_date": "2026-09-22"}, {"$set": {"status": "rejected"}})
    assert _n09(_stars(s, "three"))["final_value"] == -1  # 4 counted
    mdb.leaves.delete_one({"employee_id": t, "start_date": "2026-09-22"})
    mdb.late_requests.update_one({"employee_id": t, "date": "2026-09-29"}, {"$set": {"status": "pending"}})
    assert _n09(_stars(s, "three"))["final_value"] == 0


# ---------- override protection
def test_override_flow(s):
    eid = _eid("example")
    line_id = _n09(_stars(s, "example"))["id"]
    for bad in (-2, 1):
        r = s.put(f"{BASE_URL}/api/brsf/stars/{line_id}/override", json={"value": bad, "reason": "t"})
        assert r.status_code == 400, (bad, r.status_code, r.text)
    r = s.put(f"{BASE_URL}/api/brsf/stars/{line_id}/override", json={"value": 0, "reason": "test"})
    assert r.status_code == 200 and r.json()["final_value"] == 0
    for _ in range(3):
        assert s.post(f"{BASE_URL}/api/brsf/recalculate", json={"month": MONTH}).status_code == 200
    n = _n09(_stars(s, "example"))
    assert n["final_value"] == 0 and n["system_value"] == -1 and n["status"] == "Manually Overridden"
    assert mdb.brsf_star_lines.count_documents({"employee_id": eid, "year": 2026, "month": 9, "code": "N09"}) == 1
    assert mdb.brsf_star_lines.count_documents({"employee_id": eid, "year": 2026, "month": 9}) == 15
    r = s.post(f"{BASE_URL}/api/brsf/stars/{line_id}/reset-override")
    assert r.status_code == 200 and r.json()["final_value"] == -1
    assert _others(_stars(s, "example")) == STATE["example_others"]


def test_child_override_rejected_for_n09(s):
    n = _n09(_stars(s, "example"))
    r = s.put(f"{BASE_URL}/api/brsf/stars/{n['id']}/child-override",
              json={"key": n["system_children"][0]["key"], "value": 0})
    assert r.status_code == 400


# ---------- export
def test_export_xlsx(s):
    r = s.get(f"{BASE_URL}/api/brsf/export", params={"month": MONTH, "format": "xlsx"})
    assert r.status_code == 200
    STATE["xlsx"] = r.content
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[0]
    assert ws["S2"].value == "Late Arrivals More than 3 Per Month (-1 Monthly)"
    assert ws["R2"].value.startswith("Behavior")
    assert ws["T1"].value == "Total Positive Stars"
    assert ws["U1"].value == "Total Negative Stars"
    assert ws["V1"].value == "Total Stars"
    meta = wb["_BRSF_Metadata"]
    rows = {m[1]: m[0] for m in meta.iter_rows(min_row=2, values_only=True)}
    rx = rows[_eid("example")]
    assert ws[f"S{rx}"].value == -1
    assert ws[f"T{rx}"].value == f"=SUM(E{rx}:J{rx})"
    assert ws[f"U{rx}"].value == f"=SUM(K{rx}:S{rx})"
    assert ws[f"V{rx}"].value == f"=T{rx}+U{rx}"
    assert "K1:S1" in [str(m) for m in ws.merged_cells.ranges]


def test_export_csv(s):
    r = s.get(f"{BASE_URL}/api/brsf/export", params={"month": MONTH, "format": "csv"})
    assert r.status_code == 200
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    h = rows[0]
    assert h[18] == "Late Arrivals" and h[19:] == ["Total Positive Stars", "Total Negative Stars", "Total Stars"]
    row = next(x for x in rows if x[0] == f"{TAG}example")
    vals = list(map(float, row[4:19]))
    assert vals[14] == -1
    assert float(row[19]) == sum(vals[:6]) and float(row[20]) == sum(vals[6:])
    assert float(row[21]) == sum(vals)
    d = _stars(s, "example")
    assert float(row[21]) == d["totals"]["net_total"]


# ---------- import
def _upload(s, content, name="brsf.xlsx"):
    r = s.post(f"{BASE_URL}/api/brsf/import/preview", data={"month": MONTH},
               files={"file": (name, content)})
    assert r.status_code == 200, r.text
    STATE["batches"].append(r.json()["batch_id"])
    return r.json()


def _edit(content, eid, cells):
    wb = load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]
    rows = {m[1]: m[0] for m in wb["_BRSF_Metadata"].iter_rows(min_row=2, values_only=True)}
    for col, v in cells.items():
        ws[f"{col}{rows[eid]}"] = v
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def test_import_n09_and_p03(s):
    t = _eid("three")
    content = _edit(STATE["xlsx"], t, {"S": -1, "F": 2})
    p = _upload(s, content)
    mine = [c for c in p["changes"] if c["employee_id"] == t]
    codes = {c["code"]: c for c in mine}
    assert set(codes) == {"N09", "P03"}, p["changes"]
    assert codes["N09"]["imported"] == -1 and codes["N09"]["target"] == "override"
    assert codes["P03"]["target"] == "manual"
    r = s.post(f"{BASE_URL}/api/brsf/import/confirm", json={"batch_id": p["batch_id"]})
    assert r.status_code == 200 and not r.json()["failed"], r.text
    d = _stars(s, "three")
    n = _n09(d)
    assert n["final_value"] == -1 and n["override_value"] == -1 and n["system_value"] == 0
    assert next(l for l in d["lines"] if l["code"] == "P03")["final_value"] == 2


def test_import_invalid_n09(s):
    t = _eid("three")
    p = _upload(s, _edit(STATE["xlsx"], t, {"S": -2}))
    errs = [e for e in p["errors"] if e["code"] == "N09"]
    assert errs and "0 or -1" in errs[0]["message"]
    assert not [c for c in p["changes"] if c["code"] == "N09" and c["employee_id"] == t]


def test_import_csv_n09(s):
    r = s.get(f"{BASE_URL}/api/brsf/export", params={"month": MONTH, "format": "csv"})
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    for x in rows:
        if x[0] == f"{TAG}example":
            x[18] = "0"
    out = io.StringIO(); csv.writer(out).writerows(rows)
    p = _upload(s, out.getvalue().encode(), "brsf.csv")
    ch = [c for c in p["changes"] if c["employee_id"] == _eid("example")]
    assert [c["code"] for c in ch] == ["N09"] and ch[0]["imported"] == 0
